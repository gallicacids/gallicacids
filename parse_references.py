"""Utilities for extracting references from Word documents."""

import re
from typing import Dict, List

import pandas as pd
from docx import Document


def parse_reference(text: str) -> Dict[str, str]:
    """Parse a single reference string.

    The function attempts to extract common reference fields from a plain
    text entry. It supports leading numeric labels (e.g. ``1.``, ``[1]`` or
    ``1)``), a DOI anywhere in the string, and a four-digit year. Remaining
    sections are split by periods and mapped to authors, title and journal.
    """

    # Remove leading number like ``1.`` or ``[1]`` or ``1)``
    num = ""
    m = re.match(r"\s*(?:\[\s*)?(?P<num>\d+)(?:\s*\]|[\.).])?\s*", text)
    if m:
        num = m.group("num")
        text = text[m.end():].strip()

    # Extract DOI (pattern starting with ``10.``)
    doi = ""
    doi_match = re.search(r"(10\.\S+)", text)
    if doi_match:
        doi = doi_match.group(1).rstrip(".")
        text = text[: doi_match.start()] + text[doi_match.end() :]

    # Extract year (4-digit number starting with 19 or 20)
    year = ""
    year_match = re.search(r"\b(19|20)\d{2}\b", text)
    if year_match:
        year = year_match.group(0)
        text = text[: year_match.start()] + text[year_match.end() :]

    # Clean leftover punctuation
    text = text.strip().strip(".")

    # Split remaining by periods
    parts = [p.strip() for p in re.split(r"\.\s*", text) if p.strip()]
    authors = parts[0] if len(parts) > 0 else ""
    title = parts[1] if len(parts) > 1 else ""
    journal = parts[2] if len(parts) > 2 else ""

    return {
        "Number": num,
        "Authors": authors,
        "Title": title,
        "Journal/Publisher": journal,
        "Year": year,
        "DOI": doi,
    }


def process_docx(input_path: str, output_path: str) -> None:
    """Parse references from ``input_path`` and write a table to ``output_path``."""

    doc = Document(input_path)
    refs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    parsed: List[Dict[str, str]] = [parse_reference(ref) for ref in refs]

    df = pd.DataFrame(parsed, columns=[
        "Number",
        "Authors",
        "Title",
        "Journal/Publisher",
        "Year",
        "DOI",
    ])

    new_doc = Document()
    table = new_doc.add_table(rows=1, cols=len(df.columns))
    for i, col in enumerate(df.columns):
        table.cell(0, i).text = col

    for row in df.itertuples(index=False):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)

    new_doc.save(output_path)


if __name__ == '__main__':
    import sys

    if len(sys.argv) != 3:
        print('Usage: python parse_references.py <input.docx> <output.docx>')
        sys.exit(1)

    process_docx(sys.argv[1], sys.argv[2])
